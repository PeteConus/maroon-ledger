"""MYFER reconciliation, past Budget forecasts (open data), BP3 capital by region, ABS context."""
import csv
import glob
import re
import openpyxl
from parse_lib import SRC, all_text, rows, pick, text_between, split_line, num


# ---------- 2025-26 MYFER, Table 3 ----------
def myfer_reconciliation():
    t = text_between(all_text("MYFER-2025-26.pdf"),
                     "Table 3: Reconciliation of net operating balance, 2025-26 Budget to 2025-26 MYFER1", "Notes:")
    r = rows(t)
    labels = [
        ("start", "2025-26 Budget net operating balance"),
        ("Taxation revisions", "Taxation revisions"),
        ("Royalty and land rent revisions", "Royalty and land rent revisions"),
        ("GST revisions", "GST revisions"),
        ("Expense measures", "Expense measures"),
        ("Additional enterprise bargaining expenses", "Additional enterprise bargaining"),
        ("Other parameter adjustments", "Other parameter adjustments"),
        ("end", "2025-26 MYFER net operating balance"),
    ]
    return ["2025-26", "2026-27", "2027-28", "2028-29"], [(k, [pick(r, lab, i) for i in range(4)]) for k, lab in labels]


# ---------- past Budgets, open data (data.qld.gov.au) ----------
def _csv(budget, table_no):
    f = glob.glob(str(SRC / "opendata" / f"{budget}-budget-table-{table_no}-*.csv"))
    assert len(f) == 1, (budget, table_no, f)
    return list(csv.DictReader(open(f[0])))


def _cols(rs):
    """Map 'YYYY-YY <kind>' -> column name for the numeric columns."""
    out = {}
    for c in rs[0].keys():
        m = re.match(r"(\d{4}-\d{2}) (Budget|Projection|Estimated Actuals?|Outcome)", c)
        if m:
            out[(m.group(1), m.group(2))] = c
    return out


def _val(rs, item, col, occurrence=1):
    seen = 0
    for r in rs:
        if r["Section item"].strip().lower() == item.lower():
            seen += 1
            if seen == occurrence:
                return num(r[col])
    raise KeyError(item)


PAST_BUDGETS = {  # budget -> (operating, GGS balance sheet, NFPS balance sheet) table numbers
    "2021-22": ("10.1", "10.4", "10.6"),
    "2022-23": ("9.1", "9.4", "9.6"),
    "2023-24": ("9.1", "9.4", "9.6"),
    "2024-25": ("9.1", "9.4", "9.6"),
    "2025-26": ("8.1", "8.4", "8.6"),
}


def past_budget_forecasts():
    """For each past Budget: its Budget-year and projection-year figures."""
    out = {}
    for b, (op, gb, nb) in PAST_BUDGETS.items():
        o, g, n = _csv(b, op), _csv(b, gb), _csv(b, nb)
        oc, gc, nc = _cols(o), _cols(g), _cols(n)
        yrs = {}
        for (fy, kind), col in oc.items():
            if kind not in ("Budget", "Projection") or fy < b:
                continue
            gcol = gc.get((fy, kind)); ncol = nc.get((fy, kind))
            yrs[fy] = {
                "revenue": _val(o, "Total Revenue from Transactions", col),
                "expenses": _val(o, "Total Expenses from Transactions", col),
                "nob": _val(o, "Net operating balance", col),
                "pnfa": _val(o, "Purchases of non-financial assets", col),
                "fiscal_balance": _val(o, "Fiscal Balance", col),
                "ggs_net_debt": _val(g, "Net Debt", gcol) if gcol else None,
                "nfps_net_debt": _val(n, "Net Debt", ncol) if ncol else None,
            }
        out[b] = yrs
    return out


# ---------- BP3 Appendix C: capital purchases by entity by region ----------
SA4_BLOCKS = [
    ["Brisbane - East", "Brisbane - North", "Brisbane - South", "Brisbane - West", "Brisbane Inner City"],
    ["Cairns", "Darling Downs - Maranoa", "Central Queensland", "Gold Coast", "Ipswich"],
    ["Logan - Beaudesert", "Mackay - Isaac - Whitsunday", "Moreton Bay - North", "Moreton Bay - South", "Queensland - Outback"],
    ["Sunshine Coast", "Toowoomba", "Townsville", "Wide Bay", "TOTAL"],
]


def bp3_regions():
    t = all_text("2026-27-budget-bp3-budget-capital-statement.pdf")
    t = t[t.index("Appendix C: Capital purchases by entity by region 2026-271"):]
    t = t[: t.index("Notes\n") if "Notes\n" in t else len(t)]
    blocks = re.split(r"\n\s*Entity2\s+\$'000", t)[1:]
    assert len(blocks) == 4, len(blocks)
    data = {}
    for names, blk in zip(SA4_BLOCKS, blocks):
        blk = blk.split("Funds allocated")[0] + "Funds allocated" + blk.split("Funds allocated")[1].split("\n")[0]
        for ln in blk.splitlines():
            lab, nums = split_line(ln)
            if not nums or not lab or lab.isdigit():
                continue
            lab = re.sub(r"\d$", "", lab).strip()
            vals = [num(x) for x in nums]
            if lab.startswith("Anticipated contingency"):
                # only the Totals column carries a number
                vals = [0.0] * (len(names) - 1) + vals[-1:]
            if len(vals) != len(names):
                raise ValueError((lab, nums, names))
            for nm, v in zip(names, vals):
                data.setdefault(nm, {})[lab] = v / 1000.0  # $'000 -> $m
    return data


# ---------- ABS context ----------
def abs_gsp():
    wb = openpyxl.load_workbook(SRC / "abs" / "5220004_Annual_QLD.xlsx", read_only=True, data_only=True)
    rs = list(wb["Data1"].iter_rows(values_only=True))
    j = next(i for i, h in enumerate(rs[0]) if h and "GROSS STATE PRODUCT: Current prices" in str(h) and str(h).startswith("All sectors"))
    out = {}
    for r in rs[10:]:
        if r[0] and r[j] is not None:
            y = r[0].year
            out[f"{y-1}-{str(y)[2:]}"] = float(r[j])
    return out


def abs_erp_june():
    wb = openpyxl.load_workbook(SRC / "abs" / "310104.xlsx", read_only=True, data_only=True)
    rs = list(wb["Data1"].iter_rows(values_only=True))
    j = next(i for i, h in enumerate(rs[0]) if h and "Persons" in str(h) and "Queensland" in str(h))
    out = {}
    for r in rs[10:]:
        if r[0] and r[0].month == 6 and r[j]:
            y = r[0].year
            out[f"{y-1}-{str(y)[2:]}"] = float(r[j])  # ERP at 30 June closing the financial year
    return out


def abs_sa4_erp_2025():
    wb = openpyxl.load_workbook(SRC / "abs" / "32180DS0001_2024-25.xlsx", read_only=True, data_only=True)
    out = {}
    for r in wb["Table 3"].iter_rows(values_only=True):
        if r and isinstance(r[2], int) and isinstance(r[9], (int, float)):
            out[r[3]] = out.get(r[3], 0) + r[9]
    return out


if __name__ == "__main__":
    print(myfer_reconciliation())
    import json
    print(json.dumps(past_budget_forecasts(), indent=0)[:1500])
    reg = bp3_regions()
    print({k: round(v["Funds allocated"], 1) for k, v in reg.items()})
    print(round(sum(v["Funds allocated"] for k, v in reg.items() if k != "TOTAL"), 1))
    print(abs_gsp()); print(abs_erp_june()); s = abs_sa4_erp_2025(); print(s, sum(s.values()))
